"""
peer.py — Peer Percentile Ranking Module
Nifty 100 Financial Intelligence Platform
Sprint 3 — Screener & Health Scoring

Computes within-group percentile ranks for 10 metrics across
all 11 peer groups defined in peer_groups.xlsx.

Metrics Ranked:
    - return_on_equity_pct    (higher = better)
    - return_on_capital_pct   (higher = better)
    - net_profit_margin_pct   (higher = better)
    - debt_to_equity          (lower  = better → inverted rank)
    - free_cash_flow_cr       (higher = better)
    - net_profit_cagr_5yr     (higher = better)
    - sales_cagr_5yr          (higher = better)
    - eps_cagr_5yr            (higher = better)
    - interest_coverage       (higher = better)
    - asset_turnover          (higher = better)

Output:
    - peer_percentiles table in SQLite
    - output/peer_comparison.xlsx (11 sheets, colour-coded)

Usage:
    from src.analytics.peer import run_peer_analysis
    peer_percentiles = run_peer_analysis()
"""

import pandas as pd
import numpy as np
import sqlite3
import os
import sys

# ── Path Configuration ────────────────────────────────────────────────────────

PROJECT_PATH = r'C:\Users\VISHNU\Downloads\nifty100_project'
sys.path.append(PROJECT_PATH)
os.chdir(PROJECT_PATH)

DB_PATH       = 'data/nifty100.db'
ANALYSIS_YEAR = '2024-03'

METRIC_COLS = [
    'return_on_equity_pct',
    'return_on_capital_pct',
    'net_profit_margin_pct',
    'debt_to_equity',
    'free_cash_flow_cr',
    'net_profit_cagr_5yr',
    'sales_cagr_5yr',
    'eps_cagr_5yr',
    'interest_coverage',
    'asset_turnover',
]

# Metrics where lower value = better performance
# These get inverted percentile rank (1 - rank)
INVERT_METRICS = {'debt_to_equity'}


# ── Data Loader ───────────────────────────────────────────────────────────────

def load_peer_data():
    """
    Loads and merges peer group definitions with latest financial ratios.

    Returns:
        DataFrame: Merged peer group + ratios DataFrame
    """
    from src.etl.loader import load_all_data

    data        = load_all_data()
    peer_groups = data['peer_groups']

    conn   = sqlite3.connect(DB_PATH)
    ratios = pd.read_sql_query(
        "SELECT * FROM financial_ratios_computed", conn
    )
    conn.close()

    # Get latest year ratios
    latest = ratios[ratios['year'] == ANALYSIS_YEAR].copy()

    # Force numeric types
    for col in METRIC_COLS:
        if col in latest.columns:
            latest[col] = pd.to_numeric(latest[col], errors='coerce')

    # Merge peer groups with ratios
    peer_df = pd.merge(peer_groups, latest, on='company_id', how='left')

    return peer_df


# ── Percentile Rank Computer ──────────────────────────────────────────────────

def compute_peer_percentiles(peer_df, metric_cols=None):
    """
    Computes percentile rank for each metric within each peer group.

    For metrics in INVERT_METRICS (e.g. D/E), the rank is inverted
    so that lower values get higher percentile scores.

    For companies not in any peer group: no rows are generated.
    The engine handles this gracefully without raising errors.

    Args:
        peer_df:     Merged peer group + ratios DataFrame
        metric_cols: List of metric column names to rank.
                     Defaults to METRIC_COLS.

    Returns:
        DataFrame: Long format with columns:
                   company_id, peer_group_name, metric,
                   value, percentile_rank, year
    """
    if metric_cols is None:
        metric_cols = METRIC_COLS

    results = []

    for group_name, group in peer_df.groupby('peer_group_name'):
        for metric in metric_cols:
            if metric not in group.columns:
                continue

            # Compute percentile rank within group
            pct_rank = group[metric].rank(pct=True, na_option='keep')

            # Invert rank for metrics where lower = better
            if metric in INVERT_METRICS:
                pct_rank = 1 - pct_rank

            for _, row in group.iterrows():
                rank_val = pct_rank[row.name]
                results.append({
                    'company_id':      row['company_id'],
                    'peer_group_name': group_name,
                    'metric':          metric,
                    'value':           round(float(row[metric]), 2)
                                       if pd.notna(row[metric]) else None,
                    'percentile_rank': round(float(rank_val), 3)
                                       if pd.notna(rank_val) else None,
                    'year':            ANALYSIS_YEAR,
                })

    return pd.DataFrame(results)


# ── Database Saver ────────────────────────────────────────────────────────────

def save_to_database(peer_percentiles):
    """
    Saves peer percentile rankings to SQLite database.

    Args:
        peer_percentiles: DataFrame from compute_peer_percentiles()

    Returns:
        int: Number of rows saved
    """
    conn = sqlite3.connect(DB_PATH)
    peer_percentiles.to_sql(
        'peer_percentiles', conn, if_exists='replace', index=False
    )
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM peer_percentiles")
    count = cursor.fetchone()[0]
    conn.close()

    print(f"peer_percentiles table saved — {count} rows")
    return count


# ── Excel Generator ───────────────────────────────────────────────────────────

def generate_peer_excel(peer_df, peer_percentiles,
                        output_path='output/peer_comparison.xlsx'):
    """
    Generates peer_comparison.xlsx with one sheet per peer group.

    Colour coding:
        Green  = percentile rank >= 0.75 (top quartile)
        Yellow = percentile rank 0.25 to 0.75 (middle)
        Red    = percentile rank <= 0.25 (bottom quartile)
        Gold   = benchmark company row

    Args:
        peer_df:          Merged peer group DataFrame
        peer_percentiles: Long format percentiles DataFrame
        output_path:      Path to save Excel file
    """
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils.dataframe import dataframe_to_rows

    GREEN  = PatternFill(
        start_color='C6EFCE', end_color='C6EFCE', fill_type='solid'
    )
    YELLOW = PatternFill(
        start_color='FFEB9C', end_color='FFEB9C', fill_type='solid'
    )
    RED    = PatternFill(
        start_color='FFC7CE', end_color='FFC7CE', fill_type='solid'
    )
    GOLD   = PatternFill(
        start_color='FFD700', end_color='FFD700', fill_type='solid'
    )
    HEADER = PatternFill(
        start_color='1F4E79', end_color='1F4E79', fill_type='solid'
    )

    wb = Workbook()
    wb.remove(wb.active)

    for group_name in peer_percentiles['peer_group_name'].unique():
        ws = wb.create_sheet(title=group_name[:31])

        group_companies = peer_df[
            peer_df['peer_group_name'] == group_name
        ]['company_id'].tolist()

        benchmark_rows = peer_df[
            (peer_df['peer_group_name'] == group_name) &
            (peer_df['is_benchmark'] == 1)
        ]['company_id'].tolist()
        benchmark_co = benchmark_rows[0] if benchmark_rows else None

        # Pivot to wide format
        group_pcts = peer_percentiles[
            peer_percentiles['peer_group_name'] == group_name
        ].pivot(
            index='company_id', columns='metric', values='percentile_rank'
        )
        group_vals = peer_percentiles[
            peer_percentiles['peer_group_name'] == group_name
        ].pivot(
            index='company_id', columns='metric', values='value'
        )

        valid_metrics = [m for m in METRIC_COLS if m in group_vals.columns]

        # Build headers
        headers = (
            ['company_id'] +
            [f"{m}_value" for m in valid_metrics] +
            [f"{m}_pct_rank" for m in valid_metrics]
        )
        ws.append(headers)
        for cell in ws[1]:
            cell.fill      = HEADER
            cell.font      = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center')

        # Write company rows
        for company in group_companies:
            row = [company]
            for m in valid_metrics:
                val = group_vals.loc[company, m] \
                    if company in group_vals.index else None
                row.append(
                    round(float(val), 2) if pd.notna(val) else None
                )
            for m in valid_metrics:
                pct = group_pcts.loc[company, m] \
                    if company in group_pcts.index else None
                row.append(
                    round(float(pct), 3) if pd.notna(pct) else None
                )
            ws.append(row)

        # Colour code cells
        pct_start_col = len(valid_metrics) + 2
        for row_idx in range(2, ws.max_row + 1):
            company_name = ws.cell(row=row_idx, column=1).value

            # Gold for benchmark company
            if company_name == benchmark_co:
                for col_idx in range(1, ws.max_column + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = GOLD
                continue

            # Colour percentile rank columns
            for col_idx in range(pct_start_col, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                val  = cell.value
                if val is not None:
                    if val >= 0.75:
                        cell.fill = GREEN
                    elif val <= 0.25:
                        cell.fill = RED
                    else:
                        cell.fill = YELLOW

        # Auto width
        for col in ws.columns:
            max_len = max(
                len(str(cell.value)) if cell.value else 0
                for cell in col
            )
            ws.column_dimensions[
                col[0].column_letter
            ].width = min(max_len + 2, 20)

    wb.save(output_path)
    print(f"Saved: {output_path}")
    print(f"Sheets: {wb.sheetnames}")


# ── Master Runner ─────────────────────────────────────────────────────────────

def run_peer_analysis():
    """
    Master function — loads data, computes percentile ranks,
    saves to database and generates Excel report.

    Returns:
        DataFrame: peer_percentiles in long format

    Example:
        from src.analytics.peer import run_peer_analysis
        peer_percentiles = run_peer_analysis()
    """
    print("Loading peer group data...")
    peer_df = load_peer_data()

    print("Computing peer percentile rankings...")
    peer_percentiles = compute_peer_percentiles(peer_df)

    print(f"  Peer groups: {peer_percentiles['peer_group_name'].nunique()}")
    print(f"  Companies:   {peer_percentiles['company_id'].nunique()}")
    print(f"  Total rows:  {len(peer_percentiles)}")

    print("\nSaving to database...")
    save_to_database(peer_percentiles)

    print("\nGenerating peer_comparison.xlsx...")
    generate_peer_excel(peer_df, peer_percentiles)

    return peer_percentiles


# ── Run directly for testing ──────────────────────────────────────────────────

if __name__ == "__main__":
    peer_percentiles = run_peer_analysis()

    print("\nIT Services — ROE Rankings:")
    it_roe = peer_percentiles[
        (peer_percentiles['peer_group_name'] == 'IT Services') &
        (peer_percentiles['metric'] == 'return_on_equity_pct')
    ].sort_values('percentile_rank', ascending=False)
    print(it_roe[['company_id', 'value',
                   'percentile_rank']].to_string(index=False))
